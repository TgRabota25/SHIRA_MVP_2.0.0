import csv,hashlib,json,os,re,sqlite3,tempfile,time
from pathlib import Path
from contextlib import contextmanager
from config import data_dir,stamp
EXTENSIONS={'.mp4','.avi','.mov','.mkv','.wmv','.m4v','.webm','.mpg','.mpeg','.png','.jpg','.jpeg','.bmp','.tif','.tiff','.webp'}
STATES={'pending','running','done','review','not_found','error','cancelled','partial','accepted'}
LABELS=dict(pending='В очереди',running='Обработка',done='Распознано',review='Нужна проверка',not_found='Не найдено',error='Ошибка',cancelled='Остановлено',partial='Есть пропуски',accepted='Подтверждено')

def identity(path):
 p=Path(path).resolve();s=p.stat()
 if not p.is_file() or not s.st_size:raise ValueError('Пустой или недоступный файл')
 h=hashlib.sha256()
 with p.open('rb') as f:
  h.update(f.read(1024*1024))
  if s.st_size>1024*1024:f.seek(max(0,s.st_size-1024*1024));h.update(f.read(1024*1024))
 return hashlib.sha256((os.path.normcase(str(p))+f'|{s.st_size}|{s.st_mtime_ns}|'+h.hexdigest()).encode('utf-8')).hexdigest()

class AppLock:
 def __init__(self):self.file=None
 def acquire(self):
  self.file=(data_dir()/'app.lock').open('a+b');self.file.seek(0);self.file.write(b'0');self.file.flush();self.file.seek(0)
  try:
   if os.name=='nt':
    import msvcrt;msvcrt.locking(self.file.fileno(),msvcrt.LK_NBLCK,1)
   else:
    import fcntl;fcntl.flock(self.file,fcntl.LOCK_EX|fcntl.LOCK_NB)
  except OSError:self.file.close();self.file=None;raise RuntimeError('SHIRA уже работает с этой папкой данных')
  return self
 def close(self):
  if self.file:self.file.close();self.file=None

class Store:
 def __init__(self,path=None):
  self.path=Path(path or data_dir()/'index-v2.sqlite3');self.path.parent.mkdir(parents=True,exist_ok=True)
  with self.db() as db:
   db.executescript('''
   PRAGMA journal_mode=WAL;
   CREATE TABLE IF NOT EXISTS tasks(
    id TEXT PRIMARY KEY,path TEXT NOT NULL,name TEXT NOT NULL,mode TEXT NOT NULL,state TEXT NOT NULL,
    cfg TEXT NOT NULL,cursor REAL NOT NULL DEFAULT 0,duration REAL NOT NULL DEFAULT 0,
    number TEXT NOT NULL DEFAULT '',market TEXT NOT NULL DEFAULT '',detail TEXT NOT NULL DEFAULT '',
    evidence TEXT NOT NULL DEFAULT '',attempts INTEGER NOT NULL DEFAULT 0,failures INTEGER NOT NULL DEFAULT 0,
    created REAL NOT NULL,updated REAL NOT NULL);
   CREATE TABLE IF NOT EXISTS hits(
    id INTEGER PRIMARY KEY,task_id TEXT NOT NULL REFERENCES tasks(id),number TEXT NOT NULL,market TEXT NOT NULL,
    ts REAL NOT NULL,method TEXT NOT NULL,state TEXT NOT NULL,evidence TEXT NOT NULL DEFAULT '',raw TEXT NOT NULL DEFAULT '',
    UNIQUE(task_id,number,ts,method));
   CREATE INDEX IF NOT EXISTS ix_hits_number ON hits(number,task_id,ts);
   CREATE INDEX IF NOT EXISTS ix_tasks_state ON tasks(state,created);
   CREATE TABLE IF NOT EXISTS gaps(id INTEGER PRIMARY KEY,task_id TEXT NOT NULL,start REAL,end REAL,reason TEXT);
   CREATE TABLE IF NOT EXISTS audit(id INTEGER PRIMARY KEY,task_id TEXT NOT NULL,at REAL,action TEXT,payload TEXT);
   PRAGMA user_version=2;
   ''')
 @contextmanager
 def db(self):
  db=sqlite3.connect(self.path,timeout=30);db.row_factory=sqlite3.Row
  try:
   db.execute('PRAGMA foreign_keys=ON');db.execute('PRAGMA synchronous=FULL');yield db;db.commit()
  except BaseException:db.rollback();raise
  finally:db.close()
 def add(self,path,cfg):
  p=Path(path).resolve()
  if p.suffix.lower() not in EXTENSIONS:raise ValueError('Неподдерживаемое расширение')
  fid=identity(p);key=hashlib.sha256((fid+'|'+cfg['mode']).encode()).hexdigest();now=time.time()
  with self.db() as db:
   cur=db.execute('INSERT OR IGNORE INTO tasks(id,path,name,mode,state,cfg,created,updated) VALUES(?,?,?,?,?,?,?,?)',(key,str(p),p.name,cfg['mode'],'pending',json.dumps(dict(cfg,source_id=fid),ensure_ascii=False),now,now))
   return key,bool(cur.rowcount)
 def get(self,key):
  with self.db() as db:
   r=db.execute('SELECT * FROM tasks WHERE id=?',(key,)).fetchone();return dict(r) if r else None
 def recover(self):
  with self.db() as db:db.execute("UPDATE tasks SET state='pending',detail='Восстановлено после прерывания',updated=? WHERE state='running'",(time.time(),))
 def claim(self):
  with self.db() as db:
   db.execute('BEGIN IMMEDIATE');row=db.execute("SELECT * FROM tasks WHERE state='pending' ORDER BY updated,created LIMIT 1").fetchone()
   if row is None:return None
   db.execute("UPDATE tasks SET state='running',attempts=attempts+1,updated=? WHERE id=?",(time.time(),row['id']));return dict(row)
 def checkpoint(self,key,result):
  if result['state'] not in STATES:raise ValueError('Недопустимый статус')
  with self.db() as db:
   row=db.execute('SELECT * FROM tasks WHERE id=?',(key,)).fetchone()
   if not row:raise KeyError(key)
   for h in result.get('hits',[]):
    db.execute('INSERT OR IGNORE INTO hits(task_id,number,market,ts,method,state,evidence,raw) VALUES(?,?,?,?,?,?,?,?)',(key,h['number'],h.get('market',''),round(h['ts'],3),h['method'],h.get('state','review'),h.get('evidence',''),h.get('raw','')[:2000]))
   for start,end,reason in result.get('gaps',[]):
    db.execute('INSERT INTO gaps(task_id,start,end,reason) VALUES(?,?,?,?)',(key,start,end,reason))
   duration=result.get('duration',row['duration']);cursor=result.get('cursor',row['cursor'])
   if cursor+1e-6<row['cursor']:raise ValueError('Контрольная точка не может двигаться назад')
   state=result['state']
   if state=='done' and row['mode']=='long':
    total=db.execute('SELECT count(*) FROM hits WHERE task_id=?',(key,)).fetchone()[0]
    gaps=db.execute('SELECT count(*) FROM gaps WHERE task_id=?',(key,)).fetchone()[0]
    if gaps:state='partial'
    elif not total:state='not_found'
   db.execute('UPDATE tasks SET state=?,cursor=?,duration=?,number=?,market=?,detail=?,evidence=?,failures=0,updated=? WHERE id=?',(state,cursor,duration,result.get('number',row['number']),result.get('market',row['market']),result.get('detail',''),result.get('evidence',row['evidence']),time.time(),key))
 def failed(self,key,reason):
  with self.db() as db:
   row=db.execute('SELECT * FROM tasks WHERE id=?',(key,)).fetchone();cfg=json.loads(row['cfg']);fail=row['failures']+1
   if row['mode']=='long' and row['duration']>0 and fail>=2:
    end=min(row['duration'],row['cursor']+cfg['segment']);db.execute('INSERT INTO gaps(task_id,start,end,reason) VALUES(?,?,?,?)',(key,row['cursor'],end,str(reason)[:2000]))
    state='partial' if end>=row['duration']-.01 else 'pending'
    db.execute('UPDATE tasks SET state=?,cursor=?,failures=0,detail=?,updated=? WHERE id=?',(state,end,'Есть необработанный интервал: '+str(reason)[:1000],time.time(),key))
   else:
    state='pending' if fail<2 else 'error'
    db.execute('UPDATE tasks SET state=?,failures=?,detail=?,updated=? WHERE id=?',(state,fail,str(reason)[:2000],time.time(),key))
 def set_state(self,key,state,detail=''):
  if state not in STATES:raise ValueError(state)
  with self.db() as db:db.execute('UPDATE tasks SET state=?,detail=?,updated=? WHERE id=?',(state,detail,time.time(),key))
 def retry(self,keys,reset=False):
  with self.db() as db:
   for key in keys:
    row=db.execute('SELECT * FROM tasks WHERE id=?',(key,)).fetchone()
    if not row or row['state']=='running':continue
    restart=reset or row['mode']=='short' or row['cursor']>=row['duration']>0
    db.execute('INSERT INTO audit(task_id,at,action,payload) VALUES(?,?,?,?)',(key,time.time(),'retry_reset' if restart else 'resume',json.dumps(dict(row),ensure_ascii=False)))
    if restart:
     db.execute('DELETE FROM hits WHERE task_id=?',(key,));db.execute('DELETE FROM gaps WHERE task_id=?',(key,))
     db.execute("UPDATE tasks SET cursor=0,duration=0,number='',market='',evidence='' WHERE id=?",(key,))
    db.execute("UPDATE tasks SET state='pending',failures=0,detail='',updated=? WHERE id=?",(time.time(),key))
 def confirm(self,key,number,market,hit_id=None):
  number=number.strip()
  if not number or len(number)>200:raise ValueError('Введите номер длиной до 200 символов')
  with self.db() as db:
   row=db.execute('SELECT * FROM tasks WHERE id=?',(key,)).fetchone()
   if row['state']=='running':raise ValueError('Остановите обработку перед исправлением')
   db.execute('INSERT INTO audit(task_id,at,action,payload) VALUES(?,?,?,?)',(key,time.time(),'manual',json.dumps(dict(previous=dict(row),number=number,market=market,hit_id=hit_id),ensure_ascii=False)))
   if hit_id:
    db.execute("UPDATE hits SET number=?,market=?,state='accepted' WHERE id=? AND task_id=?",(number,market,hit_id,key))
   else:
    db.execute("UPDATE tasks SET number=?,market=?,state='accepted',detail='Проверено вручную',updated=? WHERE id=?",(number,market,time.time(),key))
 def page(self,search='',state='',mode='',sort='created',descending=True,page=0,limit=200):
  allowed={'name':'t.name COLLATE NOCASE','state':'t.state','number':'t.number','mode':'t.mode','path':'t.path COLLATE NOCASE','cursor':'t.cursor','duration':'t.duration','created':'t.created','updated':'t.updated','market':'t.market','hit_count':'(SELECT count(*) FROM hits h WHERE h.task_id=t.id)'}
  order=allowed.get(sort,allowed['created']);clauses=[];params=[]
  if search:
   needle='%'+search.replace('!','!!').replace('%','!%').replace('_','!_')+'%'
   clauses.append("(t.path LIKE ? ESCAPE '!' OR t.number LIKE ? ESCAPE '!' OR EXISTS(SELECT 1 FROM hits h WHERE h.task_id=t.id AND h.number LIKE ? ESCAPE '!'))");params.extend([needle]*3)
  if state:clauses.append('t.state=?');params.append(state)
  if mode:clauses.append('t.mode=?');params.append(mode)
  where=' WHERE '+' AND '.join(clauses) if clauses else ''
  with self.db() as db:
   total=db.execute('SELECT count(*) FROM tasks t'+where,params).fetchone()[0]
   rows=db.execute('SELECT t.*,(SELECT count(*) FROM hits h WHERE h.task_id=t.id) hit_count FROM tasks t'+where+' ORDER BY '+order+(' DESC' if descending else ' ASC')+',t.id LIMIT ? OFFSET ?',params+[limit,max(0,page)*limit]).fetchall()
   return [dict(r) for r in rows],total
 def hits(self,key=None,search='',page=0,sort='ts',descending=False):
  clauses=[];params=[]
  if key:clauses.append('h.task_id=?');params.append(key)
  if search:clauses.append('h.number LIKE ?');params.append('%'+search.replace('%','').replace('_','')+'%')
  where=' WHERE '+' AND '.join(clauses) if clauses else ''
  with self.db() as db:
   allowed={'ts':'h.ts','number':'h.number','method':'h.method','state':'h.state','name':'t.name COLLATE NOCASE','path':'t.path COLLATE NOCASE'}
   order=allowed.get(sort,'h.ts')+(' DESC' if descending else ' ASC')
   rows=db.execute('SELECT h.*,t.path,t.name FROM hits h JOIN tasks t ON t.id=h.task_id'+where+' ORDER BY '+order+',h.id LIMIT 200 OFFSET ?',params+[page*200]).fetchall()
   total=db.execute('SELECT count(*) FROM hits h'+where,params).fetchone()[0]
   return [dict(r) for r in rows],total
 def summary(self,key=None,search='',limit=200):
  clauses=[];params=[]
  if key:clauses.append('h.task_id=?');params.append(key)
  if search:clauses.append('h.number LIKE ?');params.append('%'+search.replace('%','').replace('_','')+'%')
  where=' WHERE '+' AND '.join(clauses) if clauses else ''
  with self.db() as db:
   return [dict(r) for r in db.execute('SELECT h.number,t.path,t.name,min(h.ts) first_seen,max(h.ts) last_seen,count(DISTINCT h.ts) samples FROM hits h JOIN tasks t ON t.id=h.task_id'+where+' GROUP BY h.task_id,h.number ORDER BY t.path,min(h.ts) LIMIT ?',params+[limit])]
 def gaps(self,key):
  with self.db() as db:return [dict(r) for r in db.execute('SELECT * FROM gaps WHERE task_id=? ORDER BY start',(key,))]
 def history(self,key):
  with self.db() as db:return [dict(r) for r in db.execute('SELECT * FROM audit WHERE task_id=? ORDER BY id DESC LIMIT 200',(key,))]
 def counts(self):
  with self.db() as db:return dict(db.execute('SELECT state,count(*) FROM tasks GROUP BY state'))
 def backup(self,destination):
  if Path(destination).resolve()==self.path.resolve():raise ValueError('Выберите другой файл')
  with self.db() as src:
   dst=sqlite3.connect(destination)
   try:src.backup(dst)
   finally:dst.close()
 def export(self,destination):
  dest=Path(destination);dest.parent.mkdir(parents=True,exist_ok=True)
  headers=['Файл','Путь','Режим','Статус','Номер','Маркетплейс','Время','Секунды','Источник','Подтверждение','Кадр','Примечание']
  def safe(v):
   if isinstance(v,str):
    v=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',v)[:32760]
    if v.lstrip().startswith(('=','+','-','@')):v="'"+v
   return v
  def rows():
   with self.db() as db:
    sql="SELECT t.*,h.id hid,h.number hn,h.market hm,h.ts,h.method,h.state hs,h.evidence he FROM tasks t LEFT JOIN hits h ON h.task_id=t.id AND t.mode='long' ORDER BY t.created,h.ts"
    for r in db.execute(sql):
     yield [safe(x) for x in [r['name'],r['path'],r['mode'],LABELS.get(r['state'],r['state']),r['hn'] if r['hid'] else r['number'],r['hm'] if r['hid'] else r['market'],stamp(r['ts'] or 0),r['ts'] or 0,r['method'] or '',r['hs'] or '',r['he'] or r['evidence'],r['detail']]]
  fd,temp=tempfile.mkstemp(dir=dest.parent,suffix=dest.suffix);os.close(fd)
  try:
   if dest.suffix.lower()=='.xlsx':
    from openpyxl import Workbook
    wb=Workbook(write_only=True);ws=wb.create_sheet('Индекс');ws.append(headers)
    row_number=1;sheet_number=1
    for row in rows():
     if row_number>=1048576:
      sheet_number+=1;ws=wb.create_sheet('Индекс '+str(sheet_number));ws.append(headers);row_number=1
     ws.append(row);row_number+=1
    wb.save(temp)
   else:
    with open(temp,'w',encoding='utf-8-sig',newline='') as f:
     writer=csv.writer(f,delimiter=';');writer.writerow(headers);writer.writerows(rows())
   os.replace(temp,dest)
  finally:
   if os.path.exists(temp):os.unlink(temp)
